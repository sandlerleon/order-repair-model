# -*- coding: utf-8 -*-
"""Disclose that the order-repair equation is the author's own prior model.

The ODE, the feedback function, the inversion rho = mu(I)(1-I)/I and the
saddle-node/hysteresis structure are all from Sandler (2026), a non-equilibrium
rotational-lattice study deposited on Zenodo and under review at Physics Open.
The BioEssays manuscript presented the equation as though introduced here.

Leaving that undisclosed would read as duplicate submission. Disclosing it also
answers the review comment that bistability is merely a consequence of the
chosen functional form: the form is the mean-field reduction of an explicitly
specified interacting-particle system, and kinetic Monte Carlo on that system
confirms the bistability survives fluctuations on a finite lattice.

    python disclose_lattice_origin.py
"""
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Inches, Pt
from fix_rubriq_edit import replace_in_paragraph, blocks

SRC = r"C:\Users\Leon\Downloads\BioEssays\Repair_Fidelity_Control_Parameter_BioEssays_v2.docx"

# --- 1. the model section: say where the equation comes from
OLD_MODEL = ("Positive feedback of this form is the standard route to bistability in "
             "cell-biological systems (Ferrell 2002), and the model inherits that "
             "behavior by construction rather than discovering it.")
NEW_MODEL = ("Positive feedback of this form is the standard route to bistability in "
             "cell-biological systems (Ferrell 2002), and the model inherits that "
             "behavior by construction rather than discovering it. The equation is not "
             "introduced here for the first time. It is the mean-field reduction of a "
             "non-equilibrium lattice model in which an ordered medium is continuously "
             "degraded by a flux that accelerates with local disorder and restored by an "
             "independent repair process, analyzed separately in a physical setting "
             "(Sandler 2026). That analysis establishes the bistability as a proposition "
             "for the reduced model and confirms by kinetic Monte Carlo that it survives "
             "on a finite lattice as a rate-dependent hysteresis loop, rather than being "
             "an artifact of the mean-field closure. What is proposed here is not the "
             "equation but the identification of its control parameter with repair "
             "fidelity in tissue, and what follows from that identification.")

# --- 2. the closing insight of that section, which now has a better answer
OLD_KEY = ("Bistability here is a consequence of the functional form chosen, not an "
           "empirical discovery. The question worth asking is what moves the threshold.")
NEW_KEY = ("Bistability here is inherited from the underlying lattice model, not "
           "discovered in tissue. The question worth asking is what moves the threshold.")

# --- 3. the limitation, which overstated how arbitrary the form is
OLD_LIM = ("The model is phenomenological. The form of \u03bc(I) and its parameters were "
           "chosen for qualitative consistency with a bistable, hysteretic system, not "
           "fitted to measurements.")
NEW_LIM = ("The model is phenomenological as applied to tissue. The form of \u03bc(I) is "
           "taken from the lattice model of Sandler (2026), where it follows from an "
           "explicitly specified site process; its parameters here were chosen for "
           "qualitative consistency with a bistable, hysteretic system and are not "
           "fitted to any biological measurement.")

REF = ("Sandler L (2026) Self-maintained order and hysteretic collapse in a "
       "non-equilibrium rotational lattice. Zenodo (preprint; under review). "
       "https://doi.org/10.5281/zenodo.21210708")

doc = Document(SRC)

applied = []
for old, new, tag in [(OLD_MODEL, NEW_MODEL, "model origin"),
                      (OLD_KEY, NEW_KEY, "section insight"),
                      (OLD_LIM, NEW_LIM, "limitation")]:
    for p in blocks(doc):
        if replace_in_paragraph(p, old, new):
            applied.append(tag)
            break

# --- 4. insert the reference in alphabetical position (after Lopez-Otin, before Seluanov)
paras = doc.paragraphs
ref_idx = next(i for i, p in enumerate(paras) if p.text.strip() == "References")
anchor = None
for i in range(ref_idx + 1, len(paras)):
    if paras[i].text.strip().startswith("Seluanov"):
        anchor = paras[i]
        break
if anchor is None:
    raise SystemExit("could not find the Seluanov entry to insert before")

new_p = anchor.insert_paragraph_before()
new_p.paragraph_format.line_spacing = anchor.paragraph_format.line_spacing
new_p.paragraph_format.space_after = anchor.paragraph_format.space_after
new_p.paragraph_format.left_indent = anchor.paragraph_format.left_indent
new_p.paragraph_format.first_line_indent = anchor.paragraph_format.first_line_indent
r = new_p.add_run(REF)
src = anchor.runs[0] if anchor.runs else None
r.font.size = src.font.size if src is not None and src.font.size else Pt(10)
r.font.name = src.font.name if src is not None else None
applied.append("reference inserted")

doc.save(SRC)
print("applied: %s" % ", ".join(applied))
print("references now: %d"
      % sum(1 for p in Document(SRC).paragraphs[ref_idx + 1:] if p.text.strip()))
