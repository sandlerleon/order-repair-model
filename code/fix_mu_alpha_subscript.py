# -*- coding: utf-8 -*-
"""Render mu_alpha as a real subscript, like rho_c and I_c already are.

The original build only converted rho_c, rho_damage, rho_stress and I_c, so
mu_alpha survived as a literal underscore. Left alone it sits in the same
paragraph as correctly subscripted rho_c and I_c, which reads as a typo.

A run has to be split in three to do this, which python-docx will not do, so
the run's XML element is deep-copied twice and the copies inserted after it.
Copying carries the run properties, and only the subscript flag is added.

    python fix_mu_alpha_subscript.py
"""
import copy
from docx import Document
from docx.oxml.ns import qn

SRC = r"C:\Users\Leon\Downloads\BioEssays\Repair_Fidelity_Control_Parameter_BioEssays_v2.docx"
TARGET = "\u03bc_\u03b1"          # mu, underscore, alpha
BASE, SUB = "\u03bc", "\u03b1"


def split_run(run, start, length):
    """Split one run into before / matched / after, returning the middle run."""
    el = run._element
    text = run.text
    mid = copy.deepcopy(el)
    after = copy.deepcopy(el)
    el.addnext(after)
    el.addnext(mid)

    def set_text(element, s):
        for t in element.findall(qn("w:t")):
            element.remove(t)
        t = element.makeelement(qn("w:t"), {})
        t.set(qn("xml:space"), "preserve")
        t.text = s
        element.append(t)

    set_text(el, text[:start])
    set_text(mid, text[start:start + length])
    set_text(after, text[start + length:])
    return mid


def make_subscript(element):
    rPr = element.find(qn("w:rPr"))
    if rPr is None:
        rPr = element.makeelement(qn("w:rPr"), {})
        element.insert(0, rPr)
    for tag in ("w:vertAlign",):
        for old in rPr.findall(qn(tag)):
            rPr.remove(old)
    va = rPr.makeelement(qn("w:vertAlign"), {})
    va.set(qn("w:val"), "subscript")
    rPr.append(va)


doc = Document(SRC)


def all_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


fixed = 0
for p in all_paragraphs(doc):
    if TARGET not in p.text:
        continue
    for run in list(p.runs):
        i = run.text.find(TARGET)
        while i >= 0:
            # replace "mu_alpha" with mu + subscript alpha
            mid = split_run(run, i, len(TARGET))
            for t in mid.findall(qn("w:t")):
                t.text = SUB
            make_subscript(mid)
            # the base character stays in the run before the split
            for t in run._element.findall(qn("w:t")):
                t.text = t.text + BASE
            fixed += 1
            i = run.text.find(TARGET)

doc.save(SRC)
print("mu_alpha occurrences converted to subscript: %d" % fixed)
