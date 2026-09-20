# -*- coding: utf-8 -*-
"""Repair the meaning-changing errors introduced by the Rubriq language edit.

The edit is kept: US spelling, tightened phrasing and the reordered opening are
all improvements and are left alone. What is fixed here are places where the
copyedit changed or broke the science, plus two mechanical defects.

Replacement is done character-by-character against a run-ownership map so that
existing formatting survives -- in particular the subscript runs in rho_c and
I_c, which a naive "collapse the paragraph into one run" replace would destroy.

    python fix_rubriq_edit.py
"""
import io
import os
from docx import Document

SRC = r"C:\Users\Leon\Downloads\BioEssays\Repair_Fidelity_Control_Parameter_BioEssays_v2.docx"

# (old, new, why) -- applied once each, in document order
FIXES = [
 ("as although it had one answer",
  "as though it had one answer",
  "grammar: 'as although'"),

 ("Compared with human fibroblasts, its fibroblasts require fewer oncogenic "
  "hits to transform in vitro; thus, the transformation of the whale at the "
  "single-cell level is not difficult.",
  "Compared with human fibroblasts, its fibroblasts require fewer oncogenic "
  "hits to transform in vitro, so the whale is not harder to transform at the "
  "single-cell level.",
  "the point is that the whale is NOT harder to transform; the edit obscured it"),

 ("By eliminating damaged cells more aggressively, the bowhead does not "
  "resolve Peto's paradox. It resolves it by not accumulating damage.",
  "The bowhead does not resolve Peto's paradox by eliminating damaged cells "
  "more aggressively. It resolves it by not accumulating the damage.",
  "fronted phrase reads as though the bowhead does eliminate aggressively"),

 ("must regenerate them, draw on a finite progenitor pool, or accept a "
  "shrinking and less functional one",
  "must regenerate them, drawing on a finite progenitor pool, or accept a "
  "shrinking and less functional one",
  "drawing on the pool is the cost of regenerating, not a third alternative"),

 ("the elimination strategy involves one of them to contain the other",
  "the elimination strategy engages one of them in order to contain the other",
  "grammar: 'involves X to contain Y'"),

 ("Three quantities are often used interchangeably and should not be as follows:",
  "Three quantities are often used interchangeably and should not be:",
  "'should not be as follows' inverts the sense of the list"),

 ("Repair capacity. How much repair the cell can perform?",
  "Repair capacity. How much repair the cell can perform.",
  "definition turned into a question"),

 ("Repair efficiency. what fraction of lesions is resolved?",
  "Repair efficiency. What fraction of lesions is resolved.",
  "definition turned into a question; lowercase start"),

 ("Repair fidelity. How accurately repairs the original sequence and structure?",
  "Repair fidelity. How accurately repair restores the original sequence and "
  "structure.",
  "lost its subject and verb: 'how accurately repair restores'"),

 ("The equilibrium condition is exactly inverted.",
  "The equilibrium condition inverts exactly.",
  "the condition can BE inverted exactly; it is not 'exactly inverted'"),

 ("When \u03b1 is uniformly repaired,",
  "When \u03b1 scales repair uniformly,",
  "'alpha is uniformly repaired' is meaningless"),

 ("The filled circles mark the fold changes.",
  "Filled circles mark the fold.",
  "'fold change' is a different quantity entirely in biology"),

 ("Uniform scaling is an assumption and a consequential assumption.",
  "Uniform scaling is an assumption, and a consequential one.",
  "repetition introduced by the edit"),

 ("the two folds annihilate in a cusp, \u03c1(I), becomes monotone",
  "the two folds annihilate in a cusp, \u03c1(I) becomes monotone",
  "stray comma makes rho(I) parenthetical"),

 ("and then decreases bistability entirely above",
  "and then loses bistability entirely above",
  "bistability is lost, not decreased"),

 ("The bowhead data provide a concrete candidate for CIRBP-mediated "
  "enhancement of double-strand break repair.",
  "The bowhead data supply a concrete candidate: CIRBP-mediated enhancement of "
  "double-strand-break repair.",
  "CIRBP enhancement IS the candidate, not what the candidate is for"),

 ("chosen for qualitative consistency with a bistable, hysteretic system that "
  "was not fitted to the measurements.",
  "chosen for qualitative consistency with a bistable, hysteretic system, not "
  "fitted to measurements.",
  "it is the parameters that were not fitted, not the system"),

 ("The exact scaling result in Section 6 is an algebraic property",
  "The exact scaling result above is an algebraic property",
  "the manuscript has no numbered sections; dangling cross-reference"),

 ("CIRBP is associated with several traits, such as the body size, metabolic "
  "rate, cell number, cold adaptation and immune profile,",
  "CIRBP is one of several traits that co-occur with the bowhead's body size, "
  "metabolic rate, cell number, cold adaptation and immune profile,",
  "CIRBP co-occurs with those traits; it is not 'associated with' them"),

 ("Firsanov et al. (2025) reported that enhanced repair is an important "
  "contributor to the phenotype but is not a complete explanation.",
  "Firsanov et al. (2025) present enhanced repair as an important contributor "
  "to the phenotype, not a complete explanation of it.",
  "avoid attributing a stronger negative claim to the source"),

 ("Tissue order is not yet observable.",
  "Tissue order is not yet an observable.",
  "'an observable' is the noun; the edit changed the meaning"),

 ("Is repair fidelity distinct from repair activity, and what differs between "
  "long-lived and short-lived mammals of similar size?",
  "Is repair fidelity, as distinct from repair activity, what differs between "
  "long-lived and short-lived mammals of similar size?",
  "the edit split one question into two unrelated ones"),

 ("then the bowhead repair principle loses most of its distinctive",
  "then the Bowhead Repair Principle loses most of its distinctive",
  "named principle, defined capitalised earlier"),

 ("the direct test of P5 and of the bowhead repair principle as stated.",
  "the direct test of P5 and of the Bowhead Repair Principle as stated.",
  "named principle, defined capitalised earlier"),

 ("https://doi.org/10.1158/0008-5472. CAN-08-3658",
  "https://doi.org/10.1158/0008-5472.CAN-08-3658",
  "a space was inserted into the DOI, breaking the link"),

 ("van Deursen JM (2014) The role of senescent cells in aging.",
  "van Deursen JM (2014) The role of senescent cells in ageing.",
  "reference titles are quoted verbatim; the published title is 'ageing'"),
]


def replace_in_paragraph(p, old, new):
    """Replace `old` with `new`, preserving every run's formatting.

    Each character is mapped to the run that owns it. The replacement inherits
    the formatting of the first character it replaces, and every other run keeps
    its own text and formatting -- which is what keeps the subscript runs in
    rho_c and I_c intact.
    """
    runs = p.runs
    if not runs:
        return False
    text = "".join(r.text for r in runs)
    i = text.find(old)
    if i < 0:
        return False
    owner = []
    for k, r in enumerate(runs):
        owner.extend([k] * len(r.text))
    j = i + len(old)
    new_text = text[:i] + new + text[j:]
    new_owner = owner[:i] + [owner[i]] * len(new) + owner[j:]
    for k, r in enumerate(runs):
        r.text = "".join(c for c, o in zip(new_text, new_owner) if o == k)
    return True


def blocks(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


def main():
    doc = Document(SRC)
    applied, missed = [], []
    for old, new, why in FIXES:
        done = False
        for p in blocks(doc):
            if replace_in_paragraph(p, old, new):
                done = True
                break
        (applied if done else missed).append((old, why))

    doc.save(SRC)

    log = ["APPLIED %d of %d" % (len(applied), len(FIXES)), ""]
    for old, why in applied:
        log.append("  + %-58s  %s" % (old[:58], why))
    if missed:
        log += ["", "NOT FOUND (text may already differ):"]
        for old, why in missed:
            log.append("  ! %-58s  %s" % (old[:58], why))
    io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "_fix_log.txt"),
            "w", encoding="utf-8").write("\n".join(log))
    print("applied %d of %d fixes; see _fix_log.txt" % (len(applied), len(FIXES)))
    if missed:
        print("MISSED %d" % len(missed))


if __name__ == "__main__":
    main()
