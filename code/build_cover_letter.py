# -*- coding: utf-8 -*-
"""Cover letter for the BioEssays Hypotheses submission.

Numbers come from the model output, for the same reason the manuscript's do.

    python build_cover_letter.py
"""
import json
import os
from docx import Document
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = r"C:\Users\Leon\Downloads\BioEssays"
R = json.load(open(os.path.join(HERE, "order_repair_results.json")))

ROB = R["robustness"]
SCA = R["scaling_A"]
CUSP_B = R["cusp"]["B"]
DOI_CODE = "10.5281/zenodo.22851970"
DOI_MS = "10.5281/zenodo.22851972"
REPO = "https://github.com/sandlerleon/order-repair-model"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(10)
st.paragraph_format.line_spacing = 1.15
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(1.0)


def P(text, bold=False, size=11, after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


P("Leon Sandler", size=10.5, after=0)
P("Independent Researcher", size=10.5, after=0)
P("Northbrook, Illinois 60062, United States", size=10.5, after=0)
P("sandler.leon@gmail.com | ORCID 0009-0007-4584-808X", size=10.5, after=14)

P("The Editors", size=10.5, after=0)
P("BioEssays", size=10.5, after=14)

P("Dear Editors,", after=10)

P("I am submitting the enclosed manuscript, \u201cRepair Fidelity as a Control "
  "Parameter Linking Cancer Resistance and Longevity: a bistable order\u2013repair "
  "hypothesis inspired by the bowhead whale,\u201d for consideration as a "
  "Hypotheses article.")

P("The central proposal is that DNA-repair fidelity is not simply a determinant "
  "of mutation rate but a control parameter governing whether multicellular "
  "tissue remains in a stable cooperative state. The recent demonstration that "
  "the bowhead whale resists cancer through enhanced and accurate repair rather "
  "than through additional tumour-suppressor barriers (Firsanov et al., Nature, "
  "2025) is the empirical entry point, but the hypothesis is general: it asks "
  "whether cancer resistance and age-associated tissue decline are two readouts "
  "of a single repair parameter rather than two traits that happen to co-occur.")

P("Why this may suit BioEssays", bold=True, after=4)
P("The paper is a conceptual framework with explicit predictions rather than a "
  "report of new data, and it engages a question of broad interest across cancer "
  "biology, ageing research and comparative physiology. It also makes a point of "
  "separating what is published, what the model contributes as interpretation, "
  "and what remains speculation \u2014 a distinction carried through a dedicated "
  "table and repeated in the limitations. BioEssays has published closely related "
  "conceptual work on the atavistic model of cancer, to which this hypothesis is "
  "compatible but which it deliberately does not require.")

P("What the model contributes", bold=True, after=4)
P("The order\u2013repair model is deliberately phenomenological and the manuscript "
  "says so before presenting any result: bistability follows from the assumed "
  "positive feedback rather than being discovered. What is not assumed is how the "
  "tipping point responds to repair fidelity, and this yields three results that "
  "I believe are new:")

for txt in [
    "When repair is scaled uniformly, the stress threshold is exactly "
    "proportional to the repair multiplier, and the fold position is invariant. "
    "This is algebraic, not fitted: the scaling factor cancels from the fold "
    "condition. Computing the threshold independently at 64 values confirms it to "
    "a relative deviation of %.0e." % SCA["ratio_max_rel_dev"],

    "The direction of the effect survives across %s randomly drawn parameter "
    "sets, in 100%% of cases, under three different ways of implementing repair "
    "enhancement. The qualitative claim does not depend on the functional form."
    % format(ROB["bistable_sets"], ","),

    "The magnitude does depend on implementation, by roughly twofold \u2014 and "
    "enhancement confined to constitutive repair abolishes bistability above a "
    "threshold multiplier of %.2f, converting catastrophic collapse into graded "
    "decline. This was not anticipated, and it turns a modelling detail into a "
    "measurable prediction that discriminates between mechanisms." % CUSP_B,
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(txt)
    r.font.size = Pt(11)

P("The manuscript closes with five predictions running from molecular to "
  "organismal, of which the last \u2014 that a single repair intervention should "
  "move cancer resistance and age-associated decline together \u2014 is the one "
  "on which I would ask the hypothesis to be judged. If it fails, the framework "
  "loses most of its distinctive value, and the manuscript says so.")

P("Openness and reproducibility", bold=True, after=4)
P("The model implementation, parameter set, random seed and figure generators "
  "are openly available at %s and archived at https://doi.org/%s. The manuscript "
  "is deposited at https://doi.org/%s. Every number in the paper is read "
  "programmatically from the model output rather than transcribed, and an audit "
  "script re-checks the built document against the model. All 17 references were "
  "verified against Crossref." % (REPO, DOI_CODE, DOI_MS))

P("Declarations", bold=True, after=4)
P("This manuscript is original, is not under consideration elsewhere, and has "
  "not been published previously other than as the archived preprint noted "
  "above. I am the sole author. I have no competing interests and the work "
  "received no external funding. As an unfunded independent researcher I am "
  "submitting under the standard subscription option rather than open access. "
  "Generative AI (Claude, Anthropic) was used to assist with literature "
  "synthesis, model implementation and drafting; I reviewed and edited all "
  "content, verified every reference, and take full responsibility for the "
  "manuscript, as stated in the Declarations.")

P("References are given in author\u2013year form under Wiley's free-format "
  "initial submission policy, and will be converted to the journal's style on "
  "revision if required.")

P("Thank you for considering the manuscript.", after=14)
P("Yours sincerely,", after=0)
P("Leon Sandler", after=0)

if not os.path.isdir(OUT):
    os.makedirs(OUT)
path = os.path.join(OUT, "BioEssays_Cover_Letter_Sandler.docx")
doc.save(path)
print("words : %d" % sum(len(p.text.split()) for p in doc.paragraphs))
print("saved : %s" % path)
