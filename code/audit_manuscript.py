# -*- coding: utf-8 -*-
"""Check the built manuscript against the model output and against itself.

Three classes of error this catches:
  * a number in the text that no longer matches the JSON it came from;
  * a reference in the list that nothing cites, or a citation with no entry --
    the defect that survived a journal retarget on an earlier paper;
  * a value described as being on the wrong side of a threshold named in the
    same sentence.
"""
import io, json, os, re
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
DOC = r"C:\Users\Leon\Downloads\BioEssays\Repair_Fidelity_Control_Parameter_BioEssays.docx"
R = json.load(open(os.path.join(HERE, "order_repair_results.json")))

d = Document(DOC)
paras = [p.text for p in d.paragraphs]
table_text = []                                  # kept out of `paras`: the
for t in d.tables:                               # reference section is located
    for row in t.rows:                           # by position in document order
        for c in row.cells:
            table_text.append(c.text)
TEXT = "\n".join(paras + table_text)             # numbers may legitimately
                                                 # appear inside the table

A1 = [r for r in R["modes"]["A"] if r["alpha"] == 1.0][0]
ROB, P = R["robustness"], R["parameters"]

CHECKS = [
    ("rho_c at alpha=1",      "%.2f" % A1["rho_c"]),
    ("rho_lower at alpha=1",  "%.2f" % A1["rho_lower"]),
    ("fold position I_c",     "%.2f" % A1["I_c"]),
    ("cusp alpha, mode B",    "%.2f" % R["cusp"]["B"]),
    ("mode A doubling ratio", "%.2f" % ROB["modes"]["A"]["median"]),
    ("mode B doubling ratio", "%.2f" % ROB["modes"]["B"]["median"]),
    ("mode C doubling ratio", "%.2f" % ROB["modes"]["C"]["median"]),
    ("bistable sets",         format(ROB["bistable_sets"], ",")),
    ("mu0",                   "%.2f" % P["mu0"]),
    ("beta",                  "%.0f" % P["beta"]),
    ("K",                     "%.2f" % P["K"]),
    ("h",                     "%.0f" % P["h"]),
]
print("NUMERIC CHECKS")
missing = 0
for name, val in CHECKS:
    ok = val in TEXT
    print("   %-24s %-10s %s" % (name, val, "ok" if ok else "NOT FOUND"))
    missing += 0 if ok else 1

# ------------------------------------------------------------- reference use
i = next(k for k, t in enumerate(paras) if t.strip() == "References")
body = "\n".join(paras[:i])
reflist = [t for t in paras[i + 1:] if t.strip()]

def surname(entry):
    return re.match(r"([A-Za-z\u00c0-\u024f'\u2019-]+)", entry).group(1)

print("\nREFERENCE CHECKS")
uncited = []
for entry in reflist:
    sn = surname(entry)
    m = re.search(r"\((\d{4})\)", entry)
    if not m:                                   # not a reference entry
        continue
    yr = m.group(1)
    # accept "Surname et al. YEAR", "Surname & Other YEAR", "Surname YEAR"
    if not re.search(re.escape(sn) + r"[^)]{0,60}" + yr, body):
        uncited.append("%s %s" % (sn, yr))
io.open(os.path.join(HERE, "_audit_reflist.txt"), "w", encoding="utf-8").write(
    "\n".join("%02d| %s" % (k, e) for k, e in enumerate(reflist)))
print("   entries            : %d (listed in _audit_reflist.txt)" % len(reflist))
print("   never cited in text: %s" % (", ".join(uncited) if uncited else "none"))

cited = set()
for m in re.finditer(r"\(([^()]*?\d{4}[^()]*?)\)", body):
    for part in re.split(r";", m.group(1)):
        mm = re.match(r"\s*([A-Za-z\u00c0-\u024f'\u2019-]+)", part.strip())
        yy = re.search(r"(\d{4})", part)
        if mm and yy:
            cited.add((mm.group(1), yy.group(1)))
listed = {(surname(e), re.search(r"\((\d{4})\)", e).group(1)) for e in reflist
          if re.search(r"\((\d{4})\)", e)}
orphan = sorted(c for c in cited if c not in listed)
print("   cited, not listed  : %s" % (", ".join("%s %s" % o for o in orphan) or "none"))

# --------------------------------------------------- contradiction sniff test
print("\nINTERNAL CONTRADICTION CHECKS")
issues = []
# The comparison word must be followed directly by the number it compares
# against. Without this, "1.89 ... but only 1.10 under constitutive-only
# enhancement" reads as a claim that 1.89 is under 1.10.
pat = re.compile(r"([\d.]+)\s*(?:mg|kWh|%|\$)?[^.]{0,70}?"
                 r"\b(above|below|under|over|exceeds|exceeding)\s+"
                 r"(?:the\s+|a\s+|about\s+)?([\d.]+)")
for sent in re.split(r"(?<=[.!?])\s+", TEXT):
    m = pat.search(sent)
    if not m:
        continue
    try:
        a, rel, b = float(m.group(1)), m.group(2), float(m.group(3))
    except ValueError:
        continue
    wrong = (rel in ("above", "over", "exceeds", "exceeding") and a < b) or \
            (rel in ("below", "under") and a > b)
    if wrong:
        issues.append(sent.strip()[:150])
print("   %s" % ("none found" if not issues else "%d found - see _audit_issues.txt"
                 % len(issues)))
io.open(os.path.join(HERE, "_audit_issues.txt"), "w", encoding="utf-8").write(
    "\n\n".join(issues) if issues else "none")

print("\nSUMMARY")
print("   numeric checks     : %d, missing %d" % (len(CHECKS), missing))
print("   uncited references : %d" % len(uncited))
print("   orphan citations   : %d" % len(orphan))
print("   contradictions     : %d" % len(issues))
print("   verdict            : %s"
      % ("PASS" if not (missing or uncited or orphan or issues) else "REVIEW NEEDED"))
